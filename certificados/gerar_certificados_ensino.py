"""Gera os certificados de Ensino e Pesquisa a partir de 'CERTIFICADOS ENSINO.xlsx'.

Fluxo:
1. Lê a aba "Lista" da planilha em planilhas/CERTIFICADOS ENSINO.xlsx.
   Sempre suba a planilha nova com esse nome exato, dentro da pasta "planilhas".
2. Seleciona só as pessoas com a coluna Status em branco (ainda não emitidas).
3. Preenche o modelo MODELOS/Certificado modelo.docx (tags {{NOME}}, {{TEXTO}},
   {{ESP}}, {{DE}}, {{ATE}}, {{CH}}, {{DIA_EMISSAO}}, {{MES_EMISSAO}}, {{ANO_EMISSAO}})
   para cada pessoa, com o texto certo conforme a Modalidade (Observership,
   Residente, Doutorando(a), Mini Fellowship) e a data de emissão sempre igual
   ao dia em que o script é rodado. Nome e Especialidade são corrigidos
   automaticamente (capitalização e acentos) antes de entrar no certificado.
4. Salva um .docx por pessoa em certificados_gerados/<data de hoje>/.
5. Marca na planilha o Status "EMITIDO dd/mm/aaaa" das linhas geradas, guardando antes
   uma cópia de segurança em planilhas/backup/.

Uso (duplo clique em Gerar_Certificados.bat, ou por linha de comando):
    .venv\\Scripts\\python.exe gerar_certificados_ensino.py
"""

import datetime as dt
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

import openpyxl
from docxtpl import DocxTemplate

PASTA_SCRIPT = Path(__file__).resolve().parent
PLANILHA_PADRAO = PASTA_SCRIPT / "planilhas" / "CERTIFICADOS ENSINO.xlsx"
ABA_PADRAO = "Lista"
MODELO_PADRAO = PASTA_SCRIPT / "MODELOS" / "Certificado modelo.docx"

TEXTO_POR_MODALIDADE = {
    "observership": "realizou OBSERVERSHIP - Experiência Médica na especialidade de",
    "observership (experiência multiprofissional)": (
        "realizou OBSERVERSHIP - Experiência Multiprofissional na especialidade de"
    ),
    "residente": "realizou Estágio Optativo para Residentes - Experiência Médica na especialidade de",
    "doutorando(a)": "realizou Estágio Optativo para Doutorandos - Experiência Médica na especialidade de",
    "mini fellowship": "realizou MINI FELLOWSHIP - Experiência Médica na especialidade de",
}

MESES_PT = {
    1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho",
    7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro",
}

# Grafia oficial das especialidades já usadas nesta planilha (com acentos e
# capitalização corretos). Serve para corrigir automaticamente entradas novas
# digitadas sem acento e/ou tudo em minúsculo/maiúsculo, comparando por uma
# chave "normalizada" (sem acento, minúscula) — ver normalizar_chave().
ESPECIALIDADES_CONHECIDAS = [
    "Anestesiologia", "Cabeça e Pescoço", "Cardiologia", "Cirurgia Bariátrica",
    "Cirurgia Cardiovascular", "Cirurgia do Aparelho Digestivo", "Cirurgia Geral",
    "Cirurgia Oncológica", "Cirurgia Pediátrica", "Cirurgia Plástica", "Cirurgia Torácica",
    "Cirurgia Torácica Pediátrica", "Clínica Médica", "Coloproctologia", "Dermatologia",
    "Emergência Cardiológica", "Emergência Clínica Adulta", "Emergência Clínica Adulta SUS e Convênio",
    "Endoscopia Digestiva - CPRE, Colangioscopia e Ecoendoscopia", "Gastroenterologia",
    "Gastroenterologia Pediátrica", "Ginecologia e Obstetrícia", "Infectologia Pediátrica",
    "Nefrologia", "Nefrologia Pediátrica", "Neurocirurgia", "Neurocirurgia - Cirurgia de Coluna",
    "Neurocirurgia Pediátrica", "Neurologia", "Oftalmologia", "Oncohematologia Pediátrica",
    "Oncologia Clínica", "Ortopedia e Traumatologia", "Otorrinolaringologia", "Pediatria",
    "Radiologia Pediátrica", "Reumatologia", "Terapia Intensiva", "Urgência e Emergência",
    "Urologia", "UTI Pediátrica",
]

CONECTORES_MINUSCULOS = {"de", "da", "do", "das", "dos", "e"}


def remover_acentos(texto: str) -> str:
    """Remove acentos de uma string, sem depender de bibliotecas externas."""
    forma_normalizada = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in forma_normalizada if not unicodedata.combining(c))


def normalizar_chave(texto: str) -> str:
    """Chave para comparar textos ignorando acento, maiusculas e espaços extras."""
    texto = remover_acentos(texto).lower()
    texto = re.sub(r"\s*-\s*", " - ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto


ESPECIALIDADES_POR_CHAVE = {normalizar_chave(nome): nome for nome in ESPECIALIDADES_CONHECIDAS}


def capitalizar_titulo(texto: str) -> str:
    """Capitaliza cada palavra (mantendo conectores em minúsculo), sem mexer em acentos."""
    palavras = texto.strip().split(" ")
    resultado = []
    for i, palavra in enumerate(palavras):
        if not palavra:
            continue
        if i > 0 and remover_acentos(palavra).lower() in CONECTORES_MINUSCULOS:
            resultado.append(palavra.lower())
        else:
            resultado.append(palavra[0].upper() + palavra[1:].lower())
    return " ".join(resultado)


def corrigir_especialidade(texto: str) -> str:
    """Corrige capitalização/acentos comparando com a lista de especialidades conhecidas.

    Se a especialidade já apareceu antes na planilha (mesmo sem acento/maiúscula
    errada), usa a grafia oficial. Caso contrário, só ajusta a capitalização.
    """
    canonica = ESPECIALIDADES_POR_CHAVE.get(normalizar_chave(texto))
    return canonica if canonica else capitalizar_titulo(texto)


def limpar_nome_arquivo(texto: str) -> str:
    """Prepara um texto para virar parte de um nome de arquivo (sem acentos/símbolos)."""
    texto = remover_acentos(texto.strip()).replace(" ", "_")
    texto = re.sub(r"[^A-Za-z0-9_.-]", "", texto)
    return texto.strip("._-") or "sem_identificacao"


def formatar_data(valor: object) -> str:
    """Formata uma data (datetime/date) da planilha como dd/mm/aaaa."""
    if isinstance(valor, dt.datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, dt.date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip() if valor is not None else ""


def converter_carga_horaria(valor: object) -> float | None:
    """Converte a carga horária (número ou texto como "80", "80h", "7,5") em número; None se inválida."""
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().lower().removesuffix("horas").removesuffix("h").strip().replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def formatar_carga_horaria(valor: object) -> str:
    """Formata a carga horária como número em texto (sem casas decimais quando inteira)."""
    numero = converter_carga_horaria(valor)
    if numero is None:
        return str(valor).strip() if valor is not None else ""
    if numero.is_integer():
        return str(int(numero))
    return str(numero).replace(".", ",")


def eh_email(valor: object) -> bool:
    return isinstance(valor, str) and "@" in valor


def status_em_branco(valor: object) -> bool:
    """True se o Status está vazio (inclusive célula só com espaços)."""
    return valor is None or str(valor).strip() == ""


COLUNAS_OBRIGATORIAS = [
    "Nome", "E-mail", "Modalidade", "Status", "Especialidade", "Periodo de", "Periodo Até", "Carga Horária",
]


def localizar_colunas(cabecalho: list[str]) -> dict[str, int]:
    """Acha o índice de cada coluna obrigatória, tolerando acento/maiúscula/espaço diferentes.

    Se faltar alguma, encerra com mensagem clara listando o que falta e o que existe.
    """
    por_chave = {normalizar_chave(nome): i for i, nome in enumerate(cabecalho) if nome}
    indices = {}
    faltando = []
    for coluna in COLUNAS_OBRIGATORIAS:
        i = por_chave.get(normalizar_chave(coluna))
        if i is None:
            faltando.append(coluna)
        else:
            indices[coluna] = i
    if faltando:
        raise SystemExit(colorir(
            f"ERRO: coluna(s) não encontrada(s) na planilha: {', '.join(faltando)}\n"
            f"Colunas encontradas: {', '.join(c for c in cabecalho if c)}\n"
            "Confira se o cabeçalho (linha 1) não foi renomeado ou apagado.",
            VERMELHO + NEGRITO,
        ))
    return indices


def ler_linhas_pendentes(planilha: Path, aba: str) -> tuple[list[dict], list[dict]]:
    """Lê a planilha e devolve (pendentes, já emitidos): Status em branco x Status preenchido."""
    wb = openpyxl.load_workbook(planilha, data_only=True)
    if aba not in wb.sheetnames:
        raise SystemExit(colorir(
            f"ERRO: Aba '{aba}' não encontrada em {planilha}. Abas disponíveis: {wb.sheetnames}", VERMELHO + NEGRITO
        ))
    ws = wb[aba]

    linhas = list(ws.iter_rows(min_row=1, values_only=True))
    if not linhas:
        raise SystemExit(colorir(f"ERRO: a aba '{aba}' está vazia.", VERMELHO + NEGRITO))
    cabecalho = [str(c).strip() if c else "" for c in linhas[0]]
    colunas = localizar_colunas(cabecalho)

    def indice(nome_coluna: str) -> int:
        return colunas[nome_coluna]

    idx_nome = indice("Nome")
    idx_email = indice("E-mail")
    idx_modalidade = indice("Modalidade")
    idx_status = indice("Status")
    idx_esp = indice("Especialidade")
    idx_de = indice("Periodo de")
    idx_ate = indice("Periodo Até")
    idx_ch = indice("Carga Horária")

    pendentes = []
    emitidos = []
    for numero_linha, linha in enumerate(linhas[1:], start=2):
        if linha[idx_nome] is None or str(linha[idx_nome]).strip() == "":
            continue
        valor_email = linha[idx_email]
        valor_esp = linha[idx_esp]
        # A planilha tem linhas novas com as colunas E-mail/Especialidade
        # trocadas de lugar. Detecta pelo "@" em vez de confiar na posição.
        if eh_email(valor_esp) and not eh_email(valor_email):
            valor_email, valor_esp = valor_esp, valor_email

        destino = pendentes if status_em_branco(linha[idx_status]) else emitidos
        destino.append(
            {
                "linha": numero_linha,
                "nome": str(linha[idx_nome]).strip(),
                "email": str(valor_email).strip() if valor_email else "",
                "modalidade": str(linha[idx_modalidade]).strip() if linha[idx_modalidade] else "",
                "especialidade": str(valor_esp).strip() if valor_esp else "",
                "de": linha[idx_de],
                "ate": linha[idx_ate],
                "ch": linha[idx_ch],
            }
        )
    return pendentes, emitidos


def validar_pessoa(pessoa: dict) -> list[str]:
    """Retorna a lista de problemas encontrados numa pessoa (vazia = tudo certo)."""
    problemas = []

    modalidade_normalizada = pessoa["modalidade"].strip().lower()
    if modalidade_normalizada not in TEXTO_POR_MODALIDADE:
        problemas.append(f"Modalidade '{pessoa['modalidade']}' não tem texto de certificado cadastrado")

    if not pessoa["especialidade"] or eh_email(pessoa["especialidade"]):
        problemas.append("Especialidade ausente ou inválida")

    if not eh_email(pessoa["email"]):
        problemas.append("E-mail ausente ou inválido")

    de_data = pessoa["de"] if isinstance(pessoa["de"], (dt.date, dt.datetime)) else None
    ate_data = pessoa["ate"] if isinstance(pessoa["ate"], (dt.date, dt.datetime)) else None
    if de_data is None or ate_data is None:
        problemas.append("Período de/até ausente ou em formato inválido")
    elif ate_data < de_data:
        problemas.append(f"Período até ({ate_data}) é anterior ao período de ({de_data})")

    if pessoa["ch"] in (None, ""):
        problemas.append("Carga horária ausente")
    else:
        carga = converter_carga_horaria(pessoa["ch"])
        if carga is None:
            problemas.append(f"Carga horária '{pessoa['ch']}' não é um número")
        elif carga <= 0:
            problemas.append(f"Carga horária '{pessoa['ch']}' tem que ser maior que zero")

    return problemas


def nome_arquivo_livre(pasta_saida: Path, nome_base: str) -> str:
    """Devolve um nome de arquivo que ainda não existe na pasta (acrescenta _2, _3... se preciso)."""
    candidato = f"{nome_base}.docx"
    contador = 2
    while (pasta_saida / candidato).exists():
        candidato = f"{nome_base}_{contador}.docx"
        contador += 1
    return candidato


def gerar_certificado(pessoa: dict, modelo: Path, pasta_saida: Path) -> str:
    """Renderiza o certificado de uma pessoa e devolve o nome do arquivo salvo (nunca sobrescreve)."""
    modalidade_normalizada = pessoa["modalidade"].strip().lower()
    nome_corrigido = capitalizar_titulo(pessoa["nome"])
    especialidade_corrigida = corrigir_especialidade(pessoa["especialidade"])
    hoje = dt.date.today()

    contexto = {
        "NOME": nome_corrigido,
        "TEXTO": TEXTO_POR_MODALIDADE[modalidade_normalizada],
        "ESP": especialidade_corrigida,
        "DE": formatar_data(pessoa["de"]),
        "ATE": formatar_data(pessoa["ate"]),
        "CH": formatar_carga_horaria(pessoa["ch"]),
        "DIA_EMISSAO": f"{hoje.day:02d}",
        "MES_EMISSAO": MESES_PT[hoje.month],
        "ANO_EMISSAO": str(hoje.year),
    }

    doc = DocxTemplate(str(modelo))
    doc.render(contexto)

    nome_base = f"Certificado_{limpar_nome_arquivo(nome_corrigido)}_{limpar_nome_arquivo(especialidade_corrigida)}"
    nome_arquivo = nome_arquivo_livre(pasta_saida, nome_base)
    doc.save(str(pasta_saida / nome_arquivo))
    return nome_arquivo


def chave_pessoa(pessoa: dict) -> str:
    """Identifica pessoa + especialidade + período, ignorando acento/maiúscula (para achar linhas repetidas).

    O período entra na chave para não bloquear quem volta para um novo estágio na mesma especialidade.
    """
    return "|".join([
        normalizar_chave(pessoa["nome"]),
        normalizar_chave(corrigir_especialidade(pessoa["especialidade"])),
        formatar_data(pessoa["de"]),
        formatar_data(pessoa["ate"]),
    ])


def especialidade_eh_nova(texto: str) -> bool:
    """True se a especialidade não está em ESPECIALIDADES_CONHECIDAS (acentos não podem ser corrigidos)."""
    return normalizar_chave(texto) not in ESPECIALIDADES_POR_CHAVE


def marcar_emitidos_na_planilha(planilha: Path, aba: str, linhas: list[int], texto_status: str) -> Path:
    """Escreve o texto de Status nas linhas informadas e salva a planilha.

    Antes de salvar, guarda uma cópia de segurança em planilhas/backup/. Devolve o caminho da cópia.
    """
    pasta_backup = planilha.parent / "backup"
    pasta_backup.mkdir(exist_ok=True)
    carimbo = dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    caminho_backup = pasta_backup / f"{planilha.stem}_antes_{carimbo}{planilha.suffix}"
    shutil.copy2(planilha, caminho_backup)

    wb = openpyxl.load_workbook(planilha)
    ws = wb[aba]
    cabecalho = [str(c.value).strip() if c.value else "" for c in ws[1]]
    coluna_status = localizar_colunas(cabecalho)["Status"] + 1
    for numero_linha in linhas:
        ws.cell(row=numero_linha, column=coluna_status, value=texto_status)
    wb.save(planilha)
    return caminho_backup


def montar_texto_docusign() -> tuple[str, str]:
    """Monta um Assunto/Mensagem genéricos, prontos para colar no DocuSign ao subir o lote."""
    assunto = "Certificados de Estágios Observership, Doutorandos e Residentes para assinatura"
    mensagem = (
        "Prezados, segue para assinatura certificados de Estágios Observership, "
        "Doutorandos e Residentes.\n\n"
        "Atenciosamente,\n"
        "Ensino e Pesquisa - Santa Casa de Porto Alegre"
    )
    return assunto, mensagem


# Cores ANSI para o terminal (o os.system("") ativa o suporte a cores no console do Windows).
os.system("")
VERMELHO = "\033[91m"
VERDE = "\033[92m"
AMARELO = "\033[93m"
CIANO = "\033[96m"
NEGRITO = "\033[1m"
RESET = "\033[0m"


def colorir(texto: str, cor: str) -> str:
    """Envolve o texto com o código de cor ANSI informado."""
    return f"{cor}{texto}{RESET}"


def mostrar_barra_progresso(atual: int, total: int, largura: int = 30) -> None:
    """Redesenha na mesma linha uma barra de progresso do tipo [#####-----] 50% (5/10)."""
    preenchido = int(largura * atual / total)
    barra = "#" * preenchido + "-" * (largura - preenchido)
    percentual = int(100 * atual / total)
    print(f"\r  {colorir(barra, VERDE)} {percentual:3d}% ({atual}/{total})", end="", flush=True)


def main() -> None:
    print(colorir(f"Lendo planilha: {PLANILHA_PADRAO}", CIANO))
    if not PLANILHA_PADRAO.exists():
        raise SystemExit(colorir(f"ERRO: Planilha não encontrada: {PLANILHA_PADRAO}", VERMELHO + NEGRITO))
    if not MODELO_PADRAO.exists():
        raise SystemExit(colorir(f"ERRO: Modelo não encontrado: {MODELO_PADRAO}", VERMELHO + NEGRITO))

    pendentes, emitidos = ler_linhas_pendentes(PLANILHA_PADRAO, ABA_PADRAO)
    print(f"Pessoas com Status em branco (pendentes de certificado): {len(pendentes)}")

    validos = []
    invalidos = []
    linha_emitida_por_chave = {}
    for pessoa in emitidos:
        linha_emitida_por_chave.setdefault(chave_pessoa(pessoa), pessoa["linha"])
    primeira_linha_por_chave = {}
    for pessoa in pendentes:
        problemas = validar_pessoa(pessoa)
        chave = chave_pessoa(pessoa)
        if chave in linha_emitida_por_chave:
            problemas.append(
                f"certificado já emitido na linha {linha_emitida_por_chave[chave]} (mesma pessoa, especialidade e período)"
            )
        elif chave in primeira_linha_por_chave:
            problemas.append(f"repetida (mesma pessoa, especialidade e período da linha {primeira_linha_por_chave[chave]})")
        else:
            primeira_linha_por_chave[chave] = pessoa["linha"]
        if problemas:
            invalidos.append((pessoa, problemas))
        else:
            validos.append(pessoa)

    if invalidos:
        print(colorir(
            f"\n{len(invalidos)} linha(s) IGNORADA(S) por problema nos dados (corrija na planilha e rode de novo):",
            VERMELHO + NEGRITO,
        ))
        for pessoa, problemas in invalidos:
            print(colorir(f"  - Linha {pessoa['linha']} ({pessoa['nome']}): {'; '.join(problemas)}", VERMELHO))

    if not validos:
        print(colorir("\nNenhum certificado gerado (nenhuma linha pendente válida).", AMARELO + NEGRITO))
        return

    novas = sorted({corrigir_especialidade(p["especialidade"]) for p in validos if especialidade_eh_nova(p["especialidade"])})
    if novas:
        print(colorir(
            "\nATENÇÃO: especialidade(s) nova(s), que nunca apareceram antes. Confira os ACENTOS no certificado "
            "e avise quem mantém o script para cadastrar:",
            AMARELO + NEGRITO,
        ))
        for especialidade in novas:
            print(colorir(f"  - {especialidade}", AMARELO))

    data_hoje = dt.date.today().isoformat()
    pasta_saida = PASTA_SCRIPT / "certificados_gerados" / data_hoje
    pasta_saida.mkdir(parents=True, exist_ok=True)

    print(colorir(f"\nGerando {len(validos)} certificado(s) em: {pasta_saida}", CIANO))
    gerados = []
    mostrar_barra_progresso(0, len(validos))
    for indice, pessoa in enumerate(validos, start=1):
        gerados.append(gerar_certificado(pessoa, MODELO_PADRAO, pasta_saida))
        mostrar_barra_progresso(indice, len(validos))
    print()
    for nome_arquivo in gerados:
        if nome_arquivo.removesuffix(".docx").rsplit("_", 1)[-1].isdigit():
            print(colorir(f"  OK  {nome_arquivo}  (já existia um arquivo com esse nome hoje; o antigo foi mantido)", AMARELO))
        else:
            print(colorir(f"  OK  {nome_arquivo}", VERDE))

    cor_resumo = AMARELO if invalidos else VERDE
    print(colorir(
        f"\nConcluído. {len(validos)} certificado(s) gerado(s), {len(invalidos)} ignorado(s).", cor_resumo + NEGRITO
    ))

    texto_status = f"EMITIDO {dt.date.today().strftime('%d/%m/%Y')}"
    try:
        caminho_backup = marcar_emitidos_na_planilha(
            PLANILHA_PADRAO, ABA_PADRAO, [p["linha"] for p in validos], texto_status
        )
        print(colorir(f"Planilha atualizada: Status '{texto_status}' gravado em {len(validos)} linha(s).", VERDE))
        print(f"Cópia de segurança da planilha antes da alteração: {caminho_backup}")
    except PermissionError:
        print(colorir(
            "\nERRO: os certificados foram gerados, mas NÃO foi possível gravar o Status na planilha "
            "(ela está aberta no Excel?).\n"
            f"Marque manualmente como '{texto_status}' as linhas: {', '.join(str(p['linha']) for p in validos)}\n"
            "Se não marcar, esses certificados serão gerados de novo na próxima execução.",
            VERMELHO + NEGRITO,
        ))
        erro_ao_gravar_status = True
    else:
        erro_ao_gravar_status = False

    assunto, mensagem = montar_texto_docusign()
    print("\n" + "=" * 60)
    print(colorir("Texto sugerido para colar no DocuSign ao subir o lote:", CIANO + NEGRITO))
    print("=" * 60)
    print(f"\nAssunto:\n{assunto}\n")
    print(f"Mensagem:\n{mensagem}")
    if erro_ao_gravar_status:
        sys.exit(1)


if __name__ == "__main__":
    try:
        main()
    except PermissionError as erro:
        print(colorir(
            f"\nERRO: sem permissão para abrir/salvar o arquivo: {erro.filename}\n"
            "Feche a planilha ou o certificado no Word/Excel e rode de novo.",
            VERMELHO + NEGRITO,
        ))
        sys.exit(1)
    except Exception as erro:  # noqa: BLE001 - mostra qualquer falha inesperada de forma visível
        print(colorir(f"\nERRO inesperado: {erro}", VERMELHO + NEGRITO))
        sys.exit(1)
